"""
PiPA — Painel Pedagógico
Gera o bloco `const DATA = {...};` do index.html a partir da planilha do simulado.

Uso:
    python3 atualizar_dados.py                        # usa dados/2_Simulado_2026.xlsx
    python3 atualizar_dados.py caminho/planilha.xlsx  # usa outra planilha

Saída: data_block.js — copie a linha gerada por cima da linha
`const DATA = ...;` que existe no index.html.

Fonte oficial dos dados: aba "Reports ORIGINAL" (alternativa marcada, gabarito e
pontuação oficial). A aba "Leticia" é usada apenas para o mapa disciplina -> questão.
Qualquer aba de cadastro pessoal dos alunos é ignorada de propósito.
"""
import json
import sys
from pathlib import Path

import openpyxl

PADRAO = Path(__file__).parent / "dados" / "2_Simulado_2026.xlsx"
TITULO = "2º Simulado 2026"


def main(caminho: Path) -> None:
    wb = openpyxl.load_workbook(caminho, data_only=True)

    # ---- fonte autoritativa: Reports ORIGINAL ----
    ws = wb["Reports ORIGINAL"]
    hdr = [c.value for c in ws[1]]
    idx: dict[int, dict[str, int]] = {}
    for i, h in enumerate(hdr):
        if h and str(h).startswith("Q "):
            p = str(h).split()
            idx.setdefault(int(p[1]), {})[p[2]] = i
    nq = max(idx)
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[3]]

    gab = "".join(
        sorted({str(r[idx[q]["Key"]]).strip() for r in rows if r[idx[q]["Key"]]})[0]
        for q in range(1, nq + 1)
    )

    # ---- mapa disciplina -> questoes (aba Leticia, linha 2) ----
    r2 = list(wb["Leticia"].iter_rows(min_row=2, max_row=2, values_only=True))[0]
    disc: list[dict] = []
    for col in range(7, 7 + nq):
        if r2[col]:
            disc.append({"nome": r2[col], "ini": col - 6, "fim": col - 6})
        else:
            disc[-1]["fim"] = col - 6

    # ---- alunos ----
    anomalias = []
    alunos = []
    for r in rows:
        ans = ""
        for q in range(1, nq + 1):
            v = r[idx[q]["Options"]]
            s = "" if v is None else str(v).strip()
            if s == "":
                ans += "-"            # questao em branco
            elif len(s) == 1:
                ans += s
            else:                      # marcacao dupla, ex.: "C, E"
                ans += "*"
                anomalias.append(
                    f'{str(r[3]).strip()} — Q{q}: marcou '
                    f'{"".join(ch for ch in s if ch.isalpha())}, '
                    f'gabarito {gab[q - 1]}, pontuado '
                    f'{int(float(r[idx[q]["Marks"]]))}'
                )
        corr = "".join(str(int(float(r[idx[q]["Marks"]]))) for q in range(1, nq + 1))
        declarado = int(float(r[5]))
        if corr.count("1") != declarado:
            raise SystemExit(
                f"Inconsistencia em {r[3]}: coluna Marks soma {corr.count('1')}, "
                f"mas 'Total de marcas' diz {declarado}."
            )
        alunos.append(
            {
                "id": int(r[2]) if r[2] else 0,
                "n": str(r[3]).strip(),
                "t": str(r[4]).strip(),
                "a": ans,
                "c": corr,
            }
        )

    alunos.sort(key=lambda s: (-s["c"].count("1"), s["n"]))
    data = {
        "titulo": TITULO,
        "totalQuestoes": nq,
        "gabarito": gab,
        "disciplinas": disc,
        "alunos": alunos,
    }
    saida = Path(__file__).parent / "data_block.js"
    saida.write_text(
        "const DATA = " + json.dumps(data, ensure_ascii=False, separators=(",", ":")) + ";",
        encoding="utf-8",
    )

    print(f"OK — {len(alunos)} alunos, {nq} questoes, {len(disc)} disciplinas.")
    print(f"Turmas: {sorted({a['t'] for a in alunos})}")
    print(f"Gravado em {saida}")
    if anomalias:
        print(f"\n{len(anomalias)} questao(oes) com dupla marcacao (pontuadas como zero):")
        for a in anomalias:
            print("  -", a)


if __name__ == "__main__":
    alvo = Path(sys.argv[1]) if len(sys.argv) > 1 else PADRAO
    if not alvo.exists():
        raise SystemExit(f"Planilha nao encontrada: {alvo}")
    main(alvo)
